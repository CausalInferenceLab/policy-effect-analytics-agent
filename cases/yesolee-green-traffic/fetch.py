"""에어코리아 월별 엑셀을 측정소 × 시간 패널로 정리한다.

입력: data/raw/<연도>/<연도>년 <월>월.xlsx
출력: data/panel.csv, data/panel.source.json
본분석은 처치 2곳·대조 9곳이며, 선정 근거와 추가 점검은 plan.yaml·README.md에 적는다.
"""

from __future__ import annotations

import json
import re
import sys
from datetime import UTC, datetime
from pathlib import Path

import pandas as pd

HERE = Path(__file__).resolve().parent
RAW = HERE / "data" / "raw"
OUT = HERE / "data" / "panel.csv"
ORIGIN = pd.Timestamp("2017-01-01", tz="Asia/Seoul")
TREATED = {"111124", "111125"}
CONTROL_CANDIDATES = {  # 대조군 검토 대상 13곳(제외 대상 포함). 본분석은 CONTROL_STATIONS 9곳만 사용한다.
    "111122",
    "111143",
    "111154",
    "111162",
    "111202",
    "111213",
    "111232",
    "111242",
    "111263",
    "111264",
    "111275",
    "111282",
    "111312",
}
# 핵심 비교군 4곳만 사용하는 민감도 분석의 대조군.
SENSITIVITY_CONTROL_STATIONS = (
    "111202",  # 신촌로
    "111154",  # 홍릉로
    "111232",  # 영등포로
    "111263",  # 도산대로
)
# 본분석 대조군 9곳. 한강대로·정릉로·강변북로·시흥대로는 제외한다.
CONTROL_STATIONS = (
    *SENSITIVITY_CONTROL_STATIONS,
    "111275",  # 천호대로
    "111312",  # 화랑로
    "111264",  # 강남대로
    "111242",  # 동작대로 중앙차로
    "111213",  # 공항대로
)
# 에어코리아 공개용 1시간 자료는 종료시각(Endtime) 표기: 01시 = 00:00~00:59:59.
# 관측 시작시각으로 변환하므로 1시간을 뺀다. 원자료 24시는 같은 날 23시로 변환된다.
# 운행제한 06:00~21:00은 원자료 07~21시, 변환 후 06~20시 구간에 해당한다.
# 근거: 환경부·국립환경과학원, 대기오염측정망 설치·운영지침(2016), 책자 142쪽(PDF 150쪽).
# https://www.keco.or.kr/download.do?uuid=201701170510459480#page=150
HOUR_LABEL = "end"
RAW_COLUMNS = ["지역", "망", "측정소코드", "측정소명", "측정일시", "NO2", "PM25"]
PANEL_COLUMNS = [
    "station_id",
    "station_name",
    "measurement_time",
    "datetime",
    "hour_idx",
    "hour",
    "treated",
    "no2",
    "pm25",
]


def load_plan() -> dict:
    import yaml

    return yaml.safe_load((HERE / "plan.yaml").read_text(encoding="utf-8"))


def validate_settings(controls, hour_label, plan):
    controls = {str(code) for code in controls}
    if not controls:
        raise ValueError(
            "대조군을 확정한 뒤 fetch.py의 CONTROL_STATIONS에 측정소 코드를 입력하세요."
        )
    if controls - CONTROL_CANDIDATES:
        raise ValueError(f"대조군 후보에 없는 코드: {sorted(controls - CONTROL_CANDIDATES)}")
    if hour_label not in {"end", "start"}:
        raise ValueError(
            "원자료 01~24시의 구간 의미를 확인한 뒤 HOUR_LABEL을 end 또는 start로 설정하세요."
        )
    if plan["time"]["freq"] != "hour" or plan["time"]["col"] != "hour_idx":
        raise ValueError("이 수집기는 time.col=hour_idx, time.freq=hour인 계획을 사용합니다.")
    return controls


def monthly_files(raw_dir: Path, start: int, end: int) -> list[Path]:
    first = (ORIGIN + pd.Timedelta(hours=start)).tz_localize(None).to_period("M")
    last = (ORIGIN + pd.Timedelta(hours=end)).tz_localize(None).to_period("M")
    expected = {str(month) for month in pd.period_range(first, last, freq="M")}
    found = {}
    for path in sorted(Path(raw_dir).rglob("*.xlsx")):
        match = re.fullmatch(r"(\d{4})년\s*(\d{1,2})월", path.stem)
        if not match or not 1 <= int(match[2]) <= 12:
            raise ValueError(f"월별 파일명 확인 필요: {path.name}")
        month = f"{int(match[1]):04d}-{int(match[2]):02d}"
        if month in expected:
            if month in found:
                raise ValueError(f"같은 월의 파일이 둘 이상입니다: {month}")
            found[month] = path
    missing = expected - found.keys()
    if missing:
        raise ValueError(
            f"원자료 월별 파일이 없습니다: {', '.join(sorted(missing))}. 저장 위치: {raw_dir}"
        )
    return [found[month] for month in sorted(expected)]


def read_month(path: Path, station_ids: set[str]) -> pd.DataFrame:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "openpyxl이 필요합니다. python -m pip install -r cases/yesolee-green-traffic/requirements.txt"
        ) from error

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Data"] if "Data" in workbook.sheetnames else workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, ())
        columns = {str(value).strip(): index for index, value in enumerate(header)}
        missing = set(RAW_COLUMNS) - columns.keys()
        if missing:
            raise ValueError(f"{path.name}: 필수 열이 없습니다: {', '.join(sorted(missing))}")
        records = []
        for row in rows:
            code = str(row[columns["측정소코드"]]).strip().removesuffix(".0")
            if code in station_ids:
                records.append([row[columns[name]] for name in RAW_COLUMNS])
        return pd.DataFrame(records, columns=RAW_COLUMNS)
    finally:
        workbook.close()


def normalise(raw: pd.DataFrame, controls: set[str], hour_label: str, plan: dict) -> pd.DataFrame:
    if hour_label not in {"end", "start"}:
        raise ValueError("시간 구간의 시작·종료 표기를 먼저 확인하세요.")
    raw = raw.copy()
    raw["측정소코드"] = (
        raw["측정소코드"].astype("string").str.strip().str.replace(r"\.0$", "", regex=True)
    )
    region = raw["지역"].astype("string").str.strip().str.split().str[0]
    selected = (
        region.isin(["서울", "서울특별시"])
        & raw["망"].astype("string").str.strip().eq("도로변대기")
        & raw["측정소코드"].isin(TREATED | controls)
    )
    raw = raw.loc[selected].copy()
    stamp = raw["측정일시"].astype("string").str.strip().str.replace(r"\.0$", "", regex=True)
    if not stamp.str.fullmatch(r"\d{10}").fillna(False).all():
        raise ValueError("측정일시는 YYYYMMDDHH 형식이어야 합니다.")
    hour = stamp.str[-2:].astype(int)
    if not hour.between(1, 24).all():
        raise ValueError("측정일시의 시각이 01~24 범위를 벗어났습니다.")
    clock = pd.to_datetime(stamp.str[:8], format="%Y%m%d", errors="raise") + pd.to_timedelta(
        hour, unit="h"
    )
    if hour_label == "end":
        clock -= pd.Timedelta(hours=1)
    clock = clock.dt.tz_localize("Asia/Seoul")
    panel = pd.DataFrame(
        {
            "station_id": raw["측정소코드"],
            "station_name": raw["측정소명"],
            "measurement_time": stamp,
            "datetime": clock,
            "hour_idx": ((clock - ORIGIN) / pd.Timedelta(hours=1)).astype("int64"),
            "hour": clock.dt.hour,
            "treated": raw["측정소코드"].isin(TREATED).astype(int),
        }
    )
    for source, target in [("NO2", "no2"), ("PM25", "pm25")]:
        values = pd.to_numeric(raw[source], errors="coerce")
        panel[target] = values.mask(values.isin([-999, float("inf"), float("-inf")]))
    in_window = panel["hour_idx"].between(int(plan["time"]["start"]), int(plan["time"]["end"]))
    pilot = panel["datetime"].ge(pd.Timestamp("2019-07-01", tz="Asia/Seoul")) & panel[
        "datetime"
    ].lt(pd.Timestamp("2019-12-01", tz="Asia/Seoul"))
    return panel.loc[in_window & ~pilot, PANEL_COLUMNS].reset_index(drop=True)


def build_panel(raw_dir=RAW, controls=None, hour_label=None, plan=None) -> pd.DataFrame:
    plan = load_plan() if plan is None else plan
    controls = validate_settings(
        CONTROL_STATIONS if controls is None else controls,
        HOUR_LABEL if hour_label is None else hour_label,
        plan,
    )
    hour_label = HOUR_LABEL if hour_label is None else hour_label
    paths = monthly_files(Path(raw_dir), int(plan["time"]["start"]), int(plan["time"]["end"]))
    parts = []
    for number, path in enumerate(paths, 1):
        print(f"[fetch] {number}/{len(paths)} {path.name}", flush=True)
        raw = read_month(path, TREATED | controls)
        parts.append(normalise(raw, controls, hour_label, plan))
    panel = pd.concat(parts, ignore_index=True)
    missing = (TREATED | controls) - set(panel["station_id"])
    if missing:
        raise ValueError(f"분석 기간에 관측이 없는 측정소: {', '.join(sorted(missing))}")
    duplicated = panel.duplicated(["station_id", "hour_idx"], keep=False)
    if duplicated.any():
        raise ValueError(f"측정소×시간 중복 {int(duplicated.sum())}행: 원자료를 확인하세요.")
    panel = panel.sort_values(["station_id", "hour_idx"]).reset_index(drop=True)
    panel.attrs["source_files"] = [str(path.relative_to(raw_dir)) for path in paths]
    return panel


def main() -> None:
    try:
        plan = load_plan()
        panel = build_panel(plan=plan)
        OUT.parent.mkdir(parents=True, exist_ok=True)
        temporary = OUT.with_suffix(".csv.tmp")
        panel.to_csv(temporary, index=False, encoding="utf-8")
        temporary.replace(OUT)
        source = dict(plan["data_sources"][0])
        source.update(
            rows=len(panel),
            fetched_at=datetime.now(UTC).isoformat(timespec="seconds"),
            query={
                "control_stations": list(CONTROL_STATIONS),
                "hour_label": HOUR_LABEL,
                "source_files": panel.attrs["source_files"],
            },
        )
        OUT.with_suffix(".source.json").write_text(
            json.dumps(source, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    except (ValueError, RuntimeError, OSError) as error:
        sys.exit(f"fetch.py: {error}")
    print(f"[fetch] {len(panel)}행 → {OUT.relative_to(HERE)}")


if __name__ == "__main__":
    main()
