"""역사적 정류장·노선 파일의 전체 행과 목표 ID를 감사한다."""

from __future__ import annotations

import csv
import io
import json
import math
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from .common import CASE, read_yaml, sha256, utcnow, write_json
from .deepening_history_data import MANIFEST, REQUESTS, verify_file

OUT = CASE / "results/deepening/history"
ALIASES = {
    "STDR_DE": "date", "NODE_ID": "node", "표준ID": "node",
    "STTN_NO": "ars", "ARS_ID": "ars", "STTN_NM": "name", "정류소명": "name",
    "CRDNT_X": "x", "X좌표": "x", "CRDNT_Y": "y", "Y좌표": "y",
    "STTN_TY": "type", "정류소타입": "type", "ROUTE_ID": "route", "ROUTEID": "route",
    "ROUTE_NM": "route_name", "노선명": "route_name", "ROUTE_ABRV": "route_abbreviation",
    "DSTNC": "authorized_distance", "ROUTE_TY": "route_type", "SSTTN_NM": "start_name",
    "ESTTN_NM": "end_name", "CARALC": "authorized_headway", "FIRCAR_TM": "first_raw",
    "LSTCAR_TM": "last_raw", "GROUP_NM": "company",
}
STOP_TYPES = {"0": "일반차로", "1": "중앙차로", "2": "일반중앙차로", "3": "가로변전일",
              "4": "가로변시간", "5": "마을버스", "6": "가상정류장"}


def scalar(value):
    if value is None:
        return ""
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("비유한 셀 값")
        if value.is_integer():
            return str(int(value))
    return str(value).strip()


def identifier(value, width=None):
    value = scalar(value)
    match = re.fullmatch(r"([0-9]+)([AB]?)", value)
    if not match:
        raise ValueError(f"잘못된 ID: {value!r}")
    digits, suffix = match.groups()
    if suffix and width != 5:
        raise ValueError(f"ARS 외 ID의 접미사: {value!r}")
    if width and len(digits) > width:
        raise ValueError(f"ID 길이 초과: {value!r}")
    return digits.zfill(width or len(digits)) + suffix


def date_value(value, years):
    value = scalar(value)
    if not re.fullmatch(r"[0-9]{8}", value):
        raise ValueError(f"기준일 형식 오류: {value!r}")
    parsed = datetime.strptime(value, "%Y%m%d")
    if parsed.year not in years:
        raise ValueError(f"허용 범위 밖 기준일: {value}")
    return parsed.date().isoformat()


def read_table(path):
    """Excel의 잘못된 A1 범위를 무시하고 CSV의 빈 마지막 열을 검증한다."""
    path = Path(path)
    if path.suffix == ".csv":
        blob = path.read_bytes()
        for encoding in ("utf-8-sig", "cp949"):
            try:
                text = blob.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError(f"지원하지 않는 CSV 인코딩: {path}")
        rows = list(csv.reader(io.StringIO(text), strict=True))
        info = {"sheet": "", "declared_rows": "", "declared_columns": "", "encoding": encoding}
    elif path.suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            if len(book.worksheets) != 1:
                raise ValueError(f"예상하지 않은 시트 수: {path}")
            sheet = book.active
            info = {"sheet": sheet.title, "declared_rows": sheet.max_row,
                    "declared_columns": sheet.max_column, "encoding": "xlsx"}
            sheet.reset_dimensions()
            rows = list(sheet.values)
        finally:
            book.close()
    else:
        raise ValueError(f"지원하지 않는 확장자: {path}")
    if not rows:
        raise ValueError(f"빈 파일: {path}")
    rows = [[scalar(v) for v in row] for row in rows]
    width = len(rows[0])
    # Excel은 행 끝 빈 셀을 생략하므로 헤더 너비까지 빈 값으로 채운다.
    if path.suffix == ".xlsx":
        rows = [r + [""] * (width - len(r)) for r in rows]
    if any(len(r) != width for r in rows):
        raise ValueError(f"행 너비 불일치: {path}")
    removed = 0
    while width and not rows[0][-1]:
        if any(r[-1] for r in rows[1:]):
            raise ValueError(f"이름 없는 열에 값 존재: {path}")
        rows = [r[:-1] for r in rows]
        width -= 1
        removed += 1
    if not width or any(not x for x in rows[0]) or len(set(rows[0])) != width:
        raise ValueError(f"빈/중복 헤더: {path}")
    if any(not any(r) for r in rows[1:]):
        raise ValueError(f"빈 데이터 행: {path}")
    info.update(actual_rows=len(rows), actual_columns=width, trailing_blank_columns=removed,
                header="|".join(rows[0]))
    return rows[0], rows[1:], info


def normalize(header, values, entry, years):
    if set(header) - set(ALIASES):
        raise ValueError(f"미지원 헤더: {set(header) - set(ALIASES)}")
    columns = [ALIASES[x] for x in header]
    if len(set(columns)) != len(columns):
        raise ValueError("정규화 후 중복 열")
    snapshot = re.search(r"\((20[0-9]{6})\)", entry["title"])
    output = []
    for values_row in values:
        row = dict(zip(columns, values_row, strict=True))
        # 날짜 관문을 다른 수치 해석보다 먼저 적용한다.
        row["date"] = date_value(row.get("date") or (snapshot.group(1) if snapshot else ""), years)
        row["date_basis"] = "record_STDR_DE" if "date" in columns else "filename_snapshot"
        row["source_id"] = entry["id"]
        row["kind"] = entry["kind"]
        if entry["kind"] == "route_identity":
            if not {"route", "route_name"} <= set(row):
                raise ValueError("노선 ID/이름 열 누락")
            row["route"] = identifier(row["route"], 9)
        else:
            if not {"node", "ars", "name", "x", "y"} <= set(row):
                raise ValueError("정류장 ID/ARS/좌표 열 누락")
            row["node"] = identifier(row["node"], 9)
            row["ars"] = identifier(row["ars"], 5)
            for key in ("x", "y"):
                number = float(row[key])
                if not math.isfinite(number):
                    raise ValueError("비유한 좌표")
                row[key] = number
            row["type"] = STOP_TYPES.get(row.get("type", ""), row.get("type", ""))
        output.append(row)
    return output


def target_match(rows, key, value, expected_ars=None):
    matches = [r for r in rows if r[key] == value]
    if len(matches) > 1:
        return "duplicate_key", {}
    if not matches:
        return "absent_from_catalog", {}
    found = matches[0]
    if expected_ars is not None and found["ars"] != expected_ars:
        return "ars_mismatch", found
    return "matched", found


def csv_rows(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def save_csv(name, rows):
    if not rows:
        raise ValueError(f"출력에 행이 없습니다: {name}")
    columns = list(dict.fromkeys(k for r in rows for k in r))
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / name
    with path.with_suffix(".csv.tmp").open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)
    path.with_suffix(".csv.tmp").replace(path)


def target_definitions(cfg):
    stations = csv_rows(CASE / cfg["station_registry"])
    if len({r["stop"] for r in stations}) != len(stations):
        raise ValueError("목표 정류장 ID 중복")
    roles = defaultdict(set)
    for role, name in cfg["route_targets"].items():
        for row in csv_rows(CASE / name):
            if row.get("eligible", "True") == "True":
                roles[identifier(row["route"], 9)].add(role)
    return stations, {r: ";".join(sorted(v)) for r, v in sorted(roles.items())}


def audit():
    cfg = read_yaml("deepening_identity_config.yaml")
    if cfg["years"] != [2023, 2024, 2025] or cfg["locked_outcome_years"] != [2026]:
        raise ValueError("역사 자료 연도 잠금 변경")
    policy = cfg["coordinate_policy"]
    if policy["physical_distance_enabled"] or policy["mode"] != "recorded_numeric_values_only":
        raise ValueError("물리 거리 계산은 별도 좌표계 검증이 필요합니다")
    manifest = json.loads(MANIFEST.read_text())
    if manifest["requests_sha256"] != sha256(REQUESTS):
        raise ValueError("수집 요청 해시 불일치")
    requested = {r["id"] for r in json.loads(REQUESTS.read_text())["requests"]}
    if {r["id"] for r in manifest["files"]} != requested or len(manifest["files"]) != len(requested):
        raise ValueError("수집 파일 누락/중복")
    stations, routes = target_definitions(cfg)
    profiles, snapshots, stops, route_rows = [], [], [], []
    metadata = []
    for entry in manifest["files"]:
        verify_file(entry)
        header, values, info = read_table(CASE / entry["path"])
        info = {"source_id": entry["id"], "kind": entry["kind"], "title": entry["title"],
                "bytes": entry["bytes"], "metadata_only": entry["metadata_only"], **info}
        profiles.append(info)
        if entry["metadata_only"]:
            metadata.append({"source_id": entry["id"], "header": header, "rows": values})
            continue
        frame = normalize(header, values, entry, cfg["years"])
        grouped = defaultdict(list)
        for row in frame:
            grouped[row["date"]].append(row)
        for date, records in sorted(grouped.items()):
            key = "route" if entry["kind"] == "route_identity" else "node"
            counts = Counter(r[key] for r in records)
            targets = []
            if key == "route":
                for route, role in routes.items():
                    status, found = target_match(records, "route", route)
                    targets.append({"source_id": entry["id"], "kind": entry["kind"], "date": date,
                                    "target_route": route, "target_roles": role, "match_status": status, **found})
                route_rows.extend(targets)
            else:
                for station in stations:
                    status, found = target_match(records, "node", station["stop"], station["ars"])
                    targets.append({"source_id": entry["id"], "kind": entry["kind"], "date": date,
                                    "target_node": station["stop"], "target_ars": station["ars"],
                                    "target_group": station["group"], "match_status": status, **found})
                stops.extend(targets)
            snapshots.append({"source_id": entry["id"], "kind": entry["kind"], "date": date,
                              "date_basis": records[0]["date_basis"], "rows": len(records),
                              "unique_ids": len(counts), "duplicate_id_rows": sum(n for n in counts.values() if n > 1),
                              "target_rows": len(targets),
                              **{f"target_{s}": sum(r["match_status"] == s for r in targets)
                                 for s in ("matched", "absent_from_catalog", "ars_mismatch", "duplicate_key")}})
        print(f"역사 파일 전체 행 확인: {entry['id']} {len(frame):,}행 / {len(grouped)} 기준일", flush=True)
    changes = recorded_changes(stops, policy["equality_tolerance"])
    cross = cross_source(stops, policy["equality_tolerance"])
    brackets = policy_brackets(stops, cfg["date_policy"]["policy_reference_date"], policy["equality_tolerance"])
    for name, rows in (("file_audit.csv", profiles), ("snapshot_audit.csv", snapshots),
                       ("target_stop_history.csv", stops), ("target_route_history.csv", route_rows),
                       ("stop_record_changes.csv", changes), ("cross_source_checks.csv", cross),
                       ("policy_bracketing_records.csv", brackets)):
        save_csv(name, rows)
    write_json(OUT / "metadata_definitions.json", metadata)
    names = ["deepening_identity_config.yaml", "gangnam/deepening_identity.py", "deepening-history-sources.json",
             cfg["station_registry"], *cfg["route_targets"].values()]
    write_json(OUT / "identity_audit.json", {
        "created_at": utcnow(), "inputs": {n: sha256(CASE / n) for n in names},
        "files": len(profiles), "data_files": sum(not r["metadata_only"] for r in profiles),
        "data_rows": sum(r["actual_rows"] - 1 for r in profiles if not r["metadata_only"]),
        "snapshots": len(snapshots), "target_stations": len(stations), "target_routes": len(routes),
        "coordinate_policy": policy, "interpolation": False,
        "interpretation": "공개 목록의 ID·좌표 기록을 확인했다. 실제 시설·검지 경계·변경 시행일의 동일성은 미확인이다.",
    })


def recorded_changes(stops, tolerance):
    """좌표 숫자·이름·유형의 기록 변경만 비교하고 관측 공백을 그대로 남긴다."""
    grouped = defaultdict(list)
    for row in stops:
        grouped[(row["kind"], row["target_node"])].append(row)
    changes = []
    for (kind, node), rows in sorted(grouped.items()):
        rows = sorted(rows, key=lambda r: (r["date"], r["source_id"]))
        for before, after in zip(rows[:-1], rows[1:], strict=True):
            valid = before["match_status"] == after["match_status"] == "matched"
            result = {"kind": kind, "target_node": node, "target_ars": after["target_ars"],
                      "target_group": after["target_group"], "from_date": before["date"], "to_date": after["date"],
                      "from_source": before["source_id"], "to_source": after["source_id"],
                      "from_status": before["match_status"], "to_status": after["match_status"],
                      "comparable": valid,
                      "days_between_records": (datetime.fromisoformat(after["date"]) - datetime.fromisoformat(before["date"])).days}
            if valid:
                result.update({f"{which}_{k}": row[k] for which, row in (("before", before), ("after", after))
                               for k in ("x", "y", "name", "type")})
                result.update(coordinate_record_changed=max(abs(after[k] - before[k]) for k in ("x", "y")) > tolerance,
                              name_changed=after["name"] != before["name"], type_changed=after["type"] != before["type"])
            changes.append(result)
    return changes


def cross_source(stops, tolerance):
    position = {(r["date"], r["target_node"]): r for r in stops if r["kind"] == "stop_positions"}
    result = []
    for median in stops:
        if median["kind"] != "median_stops":
            continue
        other = position.get((median["date"], median["target_node"]))
        valid = bool(other and other["match_status"] == median["match_status"] == "matched")
        item = {"date": median["date"], "target_node": median["target_node"], "target_ars": median["target_ars"],
                "target_group": median["target_group"], "median_status": median["match_status"],
                "position_status": other["match_status"] if other else "no_same_date_file", "comparable": valid}
        if valid:
            item.update(coordinates_agree=max(abs(median[k] - other[k]) for k in ("x", "y")) <= tolerance,
                        name_agrees=median["name"] == other["name"])
        result.append(item)
    return result


def policy_brackets(stops, policy_date, tolerance):
    """서로 다른 목록일 수 있는 최근 전·첫 후 기록을 명시적으로 비교한다."""
    groups = defaultdict(list)
    for row in stops:
        groups[row["target_node"]].append(row)
    output = []
    for node, rows in sorted(groups.items()):
        valid = [r for r in rows if r["match_status"] == "matched"]
        before = sorted((r for r in valid if r["date"] < policy_date),
                        key=lambda r: (r["date"], r["kind"] == "stop_positions"))
        after = sorted((r for r in valid if r["date"] >= policy_date),
                       key=lambda r: (r["date"], r["kind"] != "stop_positions"))
        result = {"target_node": node, "target_ars": rows[0]["target_ars"],
                  "target_group": rows[0]["target_group"], "policy_reference_date": policy_date,
                  "has_both_records": bool(before and after), "physical_change_date_verified": False}
        if before and after:
            pre, post = before[-1], after[0]
            result.update({f"{which}_{k}": row[k] for which, row in (("before", pre), ("after", post))
                           for k in ("source_id", "kind", "date", "x", "y", "name", "type")})
            result.update(coordinate_record_changed=max(abs(post[k] - pre[k]) for k in ("x", "y")) > tolerance,
                          name_changed=post["name"] != pre["name"],
                          days_between_records=(datetime.fromisoformat(post["date"]) - datetime.fromisoformat(pre["date"])).days,
                          same_catalog_kind=pre["kind"] == post["kind"])
        output.append(result)
    return output
