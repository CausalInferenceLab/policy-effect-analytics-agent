"""TOPIS 월별 XLSX를 검증하여 보관하고 선택 도로만 순차 파싱한다."""

from __future__ import annotations

import calendar
import json
import math
from datetime import timedelta
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import pandas as pd
import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

from .common import digest, read_yaml, sha256, utcnow, write_json

BASE = "https://topis.seoul.go.kr"
SOURCE = BASE + "/refRoom/openRefRoom_1.do"
META = ["일자", "요일", "도로명", "링크아이디", "시점명", "종점명", "방향", "거리",
        "차선수", "기능유형구분", "도심/외곽구분", "권역구분"]
HEADERS = META + [f"~{i:02d}시" for i in range(1, 25)]
LEGACY_HEADERS = META + [f"{i:02d}시" for i in range(1, 25)]


def session():
    s = requests.Session()
    # 조회용 POST만 사용한다. 게시·수정 요청은 없다.
    retry = Retry(total=3, backoff_factor=1, status_forcelist=[429, 500, 502, 503, 504],
                  allowed_methods={"GET", "POST"})
    s.mount("https://", HTTPAdapter(max_retries=retry))
    s.headers.update({"Referer": SOURCE, "User-Agent": "public-data-course-gangnam/1.0"})
    return s


def list_months(s, year):
    with s.post(BASE + "/refroom/selectRefRoomListASC.do",
                data={"blbdDivCd": "06", "bdwrDivCd": str(year), "mainBdwrRowNum": "12"},
                timeout=(15, 90)) as r:
        r.raise_for_status()
        rows = r.json()["rows"]
    out = {}
    for row in rows:
        month = int(row["months"])
        if month in out or not 1 <= month <= 12 or str(row["bdwrDivCd"]) != str(year):
            raise ValueError("TOPIS 목록의 월·연도·중복을 확인하세요")
        out[month] = {k: row[k] for k in ("apndFileNm", "apndFilePathNm", "bdwrSeq", "blbdDivCd")}
    return out


def validate_xlsx(path):
    try:
        with ZipFile(path) as z:
            if "xl/workbook.xml" not in z.namelist() or z.testzip() is not None:
                raise ValueError("XLSX 구조 또는 CRC 오류")
    except BadZipFile as e:
        raise ValueError("XLSX 대신 HTML 오류 응답 또는 손상된 ZIP을 받았습니다") from e


def download(s, entry, month, folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{month}.xlsx"
    manifest_path = folder / f"{month}.source.json"
    if path.exists() and manifest_path.exists():
        m = json.loads(manifest_path.read_text(encoding="utf-8"))
        if m["request"] == entry and m["sha256"] == sha256(path):
            validate_xlsx(path)
            return path, m
        raise ValueError(f"캐시 해시 또는 원격 항목 변경: {path}. 기존 자료를 보존하고 확인하세요")
    if path.exists():
        raise ValueError(f"매니페스트 없는 기존 파일을 덮어쓰지 않습니다: {path}")
    temp = path.with_suffix(".xlsx.tmp")
    try:
        with s.post(BASE + "/downloadFileRefRoom.do", data=entry, stream=True,
                    timeout=(15, 120)) as r:
            r.raise_for_status()
            if "html" in r.headers.get("Content-Type", "").lower():
                raise ValueError("다운로드가 HTML 오류 페이지를 반환했습니다")
            with temp.open("wb") as f:
                for chunk in r.iter_content(1024 * 1024):
                    f.write(chunk)
        validate_xlsx(temp)
        manifest = {"month": month, "source": SOURCE, "url": BASE + "/downloadFileRefRoom.do",
                    "request": entry, "retrieved_at": utcnow(), "bytes": temp.stat().st_size,
                    "sha256": sha256(temp)}
        temp.replace(path)
        write_json(manifest_path, manifest)
        return path, manifest
    finally:
        temp.unlink(missing_ok=True)


def interval_start(day, hour_label):
    """공식 종료시각 표기 확인 후 사용. ~24시도 원래 날짜의 23시다."""
    if hour_label not in HEADERS[12:]:
        raise ValueError(f"알 수 없는 시간 열: {hour_label}")
    raw = str(day).strip()
    if raw.endswith(".0"):
        raw = raw[:-2]
    date = pd.to_datetime(raw, format="%Y%m%d", errors="raise")
    return date.tz_localize("Asia/Seoul") + timedelta(hours=int(hour_label[1:3]) - 1)


def speed_value(raw, rules):
    if raw is None or str(raw).strip() == "":
        if not rules["blank_is_missing"]:
            raise ValueError("빈 셀의 처리 정의가 없습니다")
        return math.nan, "blank"
    if str(raw).strip() in {str(x) for x in rules["sentinels"]}:
        return math.nan, "sentinel"
    try:
        value = float(raw)
    except (TypeError, ValueError) as e:
        raise ValueError(f"정의되지 않은 속도 값: {raw!r}") from e
    if not math.isfinite(value):
        raise ValueError("정의되지 않은 비유한 속도 값")
    for marker in rules["sentinels"]:
        try:
            numeric_marker = float(marker)
        except (TypeError, ValueError):
            continue
        if value == numeric_marker:
            return math.nan, "sentinel"
    if value < 0:
        if rules["negative"] != "missing":
            raise ValueError("음수 정의 미확인")
        return math.nan, "negative"
    if value == 0:
        if rules["zero"] == "missing":
            return math.nan, "zero"
        if rules["zero"] != "valid":
            raise ValueError("0 속도 정의 미확인")
    return value, ""


def parse_month(path, month, roads, rules=None):
    """rules=None이면 속도값에 접근하지 않고 메타데이터만 추출한다."""
    from openpyxl import load_workbook

    meta, records, seen, first = {}, [], set(), None
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = list(next(rows))
        if header not in (HEADERS, LEGACY_HEADERS):
            raise ValueError(f"헤더 변경 감지: {header}")
        for row in rows:
            if all(v is None for v in row):
                continue
            if len(row) != len(HEADERS):
                raise ValueError("열 개수 불일치")
            if row[2] not in roads:
                continue
            day = interval_start(row[0], "~01시")
            if day.strftime("%Y-%m") != month:
                raise ValueError(f"요청 월 밖의 자료: {day}")
            if row[3] is None or str(row[3]).strip() == "":
                raise ValueError("링크 ID가 비어 있습니다")
            link = str(row[3]).removesuffix(".0")
            key = (link, str(day.date()))
            if key in seen:
                raise ValueError(f"링크×일자 중복: {key}")
            seen.add(key)
            m = dict(zip(META[2:], row[2:12], strict=True))
            m["링크아이디"] = link
            variant = digest(m)
            meta.setdefault(variant, {**m, "observed_from": str(day.date()),
                                     "observed_to": str(day.date()), "days": 0})
            meta[variant]["observed_to"] = max(meta[variant]["observed_to"], str(day.date()))
            meta[variant]["observed_from"] = min(meta[variant]["observed_from"], str(day.date()))
            meta[variant]["days"] += 1
            first = day if first is None else min(first, day)
            if rules is not None:
                for i, raw in enumerate(row[12:]):
                    value, reason = speed_value(raw, rules)
                    records.append((link, day + timedelta(hours=i), value, not reason, reason))
        if not seen:
            raise ValueError(f"선택 도로의 자료가 없습니다: {month}")
    finally:
        wb.close()
    expected_days = calendar.monthrange(int(month[:4]), int(month[5:]))[1]
    summary = {"month": month, "header": header, "selected_link_days": len(seen),
               "expected_days": expected_days, "metadata": sorted(meta.values(), key=lambda x: (x["링크아이디"], x["observed_from"]))}
    frame = pd.DataFrame(records, columns=["link_id", "interval_start_kst", "speed_kmh", "is_valid", "invalid_reason"])
    return frame, summary


def check_metadata(summaries):
    """선택 도로의 링크 속성 변경을 감지하여 같은 링크로 조용히 합치지 않는다."""
    by_link = {}
    for summary in summaries:
        for row in summary["metadata"]:
            attributes = {k: row[k] for k in META[2:]}
            by_link.setdefault(row["링크아이디"], set()).add(digest(attributes))
    changed = sorted(link for link, signatures in by_link.items() if len(signatures) > 1)
    if changed:
        raise ValueError(f"메타데이터 변경 링크를 확인하고 범위를 다시 고정하세요: {changed}")


def check_corridors(summaries, cfg):
    for name, spec in cfg["corridors"].items():
        if spec["status"] != "verified":
            continue
        links = spec["links"]
        if not links or len({x["id"] for x in links}) != len(links):
            raise ValueError(f"고정 링크 목록 오류: {name}")
        if any(a["end"] != b["start"] for a, b in zip(links, links[1:], strict=False)):
            raise ValueError(f"시종점 이름이 연결되지 않는 도로 구간: {name}")
        for summary in summaries:
            observed = {m["링크아이디"]: m for m in summary["metadata"]}
            for link in links:
                row = observed.get(link["id"])
                if row is None or (row["시점명"], row["종점명"], row["방향"], float(row["거리"])) != (
                        link["start"], link["end"], spec["direction"], float(link["length"])):
                    raise ValueError(f"계획과 실제 링크 속성 불일치: {summary['month']} {name} {link['id']}")


def read_hourly(folder, verify=False):
    paths = sorted(Path(folder).glob("*.parquet"))
    if not paths:
        raise ValueError("정규화된 월별 Parquet가 없습니다. fetch.py를 먼저 실행하세요")
    if verify:
        contract = read_yaml("data_contract.yaml")
        roads = read_yaml("corridors.yaml")["candidate_roads"]
        for path in paths:
            info = json.loads(path.with_suffix(".json").read_text(encoding="utf-8"))
            if (info["contract_hash"] != digest(contract) or info["road_selection_hash"] != digest(roads)
                    or info["parquet_sha256"] != sha256(path)):
                raise ValueError(f"설정 변경 또는 가공 파일 무결성 오류: {path}")
        summaries = [json.loads((Path(folder).parent / "metadata" / f"{p.stem}.json").read_text(encoding="utf-8")) for p in paths]
        check_metadata(summaries)
        check_corridors(summaries, read_yaml("corridors.yaml"))
    frame = pd.concat([pd.read_parquet(p) for p in paths], ignore_index=True)
    if frame.duplicated(["link_id", "interval_start_kst"]).any():
        raise ValueError("월별 자료 사이에 링크×시간 중복이 있습니다")
    return frame.sort_values(["interval_start_kst", "link_id"])
