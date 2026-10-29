"""경기 대표일 XLSX를 전체 XML 범위로 읽고 원래 키·값을 별도 Parquet에 보존한다."""

import io
import json
import re
import shutil
import zipfile
from collections import Counter

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

from .common import CASE, read_yaml, sha256, utcnow, write_json
from .deepening_identity import scalar

CONFIG = CASE / "deepening_gyeonggi_audit_config.yaml"
RAW = CASE / "data/deepening/gyeonggi/processed"
OUT = CASE / "results/deepening/gyeonggi"
COMMON = ["시/군/구", "연도", "월", "노선", "시종점", "일시"]
TRIP_FIELDS = [f"{kind}통행_{action}" for kind in ["단일", "환승", "전체"] for action in ["승차", "하차", "전체"]]
SCHEMA = pa.schema([(k, pa.string()) for k in ["source", "member", "family", "city", "route", "terminal", "day_type", "detail", "policy_name"]]
                   + [(k, pa.int64()) for k in ["year", "month", "hour", *[f"v{i}" for i in range(9)]]])


def family_for(header):
    if header == COMMON + ["시간", "이용객수"]:
        return "hourly"
    if header == COMMON + ["상세", "이용객수"]:
        return "daily"
    if header == COMMON + TRIP_FIELDS:
        return "trips"
    raise ValueError(f"unrecognized workbook schema: {header}")


def integer(value, missing=False):
    text = scalar(value)
    if missing and not text:
        return None
    if not re.fullmatch(r"-?\d+", text):
        raise ValueError(f"non-integer value: {text!r}")
    return int(text)


def aliases(cfg):
    result = {n: n for n in cfg["policy_names"]}
    for name in cfg["explicit_ab_variants"]:
        for suffix in "AB":
            if name + suffix in result and result[name + suffix] != name:
                raise ValueError("ambiguous name candidate")
            result[name + suffix] = name
    return result


def normalize(row, family, expected_year, cfg):
    width = 15 if family == "trips" else 8
    if len(row) > width and any(v is not None for v in row[width:]):
        raise ValueError("nonempty workbook columns beyond header")
    row = list(row[:width]) + [None] * max(0, width - len(row))
    year = integer(row[1])
    if year != expected_year or year not in cfg["years"] or year >= cfg["locked_year"]:
        raise ValueError("unexpected or locked source year")
    month = integer(row[2])
    city, route, terminal, day = (scalar(row[i]) for i in [0, 3, 4, 5])
    if month not in range(1, 13) or not city or not route or day not in {"평일", "주말"}:
        raise ValueError("invalid month, city, route or day type")
    detail = scalar(row[6]) if family != "trips" else "전일"
    hour = -1
    if family == "hourly":
        match = re.fullmatch(r"(\d{2})시~(\d{2})시", detail)
        if not match or int(match[1]) not in range(24) or int(match[2]) % 24 != (int(match[1]) + 1) % 24:
            raise ValueError(f"unrecognized hour interval: {detail}")
        hour = int(match[1])
    elif detail not in {"전일", "주간"}:
        raise ValueError(f"unrecognized daily detail: {detail}")
    values = [integer(x, missing=True) for x in (row[6:] if family == "trips" else row[7:])]
    return dict(family=family, city=city, route=route, terminal=terminal, day_type=day,
                detail=detail, year=year, month=month, hour=hour,
                **{f"v{i}": values[i] if i < len(values) else None for i in range(9)})


def parse_book(source, member, handle, cfg):
    book = openpyxl.load_workbook(handle, read_only=True, data_only=True)
    writer = None
    try:
        if set(book.sheetnames) != {"워크시트 익스포트", "SQL"}:
            raise ValueError("unexpected data or SQL sheets")
        sheet = book["워크시트 익스포트"]
        info = dict(source=source["id"], member=member, year=source["year"],
                    declared_rows=sheet.max_row, declared_columns=sheet.max_column)
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        header = [scalar(x) for x in next(rows)]
        family = family_for(header)
        path = RAW / f"{source['id']}_{family}.parquet"
        part = path.with_suffix(".part")
        sidecar = path.with_suffix(".audit.json")
        identity = dict(configuration_sha256=sha256(CONFIG), parser_sha256=sha256(__file__))
        if path.exists() and sidecar.exists():
            saved = json.loads(sidecar.read_text())
            if (all(saved[k] == v for k, v in identity.items()) and saved["source_sha256"] == source["sha256"]
                    and saved["member"] == member and sha256(path) == saved["sha256"]
                    and sha256(CASE / saved["sql_path"]) == saved["sql_sha256"]):
                print(f"verified cached workbook: {source['id']} {family}", flush=True)
                return saved
            raise ValueError("existing workbook extraction has different inputs")
        if path.exists() or part.exists():
            raise ValueError("parser destination exists without a verified complete manifest")
        writer = pq.ParquetWriter(part, SCHEMA, compression="zstd")
        counter, batch, names = Counter(), [], aliases(cfg)
        for row in rows:
            item = normalize(row, family, source["year"], cfg)
            item.update(source=source["id"], member=member, policy_name=names.get(item["route"], ""))
            vals = [item[f"v{i}"] for i in range(9 if family == "trips" else 1)]
            counter.update(rows=1, missing_values=sum(v is None for v in vals),
                           zero_values=sum(v == 0 for v in vals), negative_values=sum(v is not None and v < 0 for v in vals),
                           name_candidate_rows=bool(item["policy_name"]))
            batch.append(item)
            if len(batch) == cfg["batch_rows"]:
                writer.write_table(pa.Table.from_pylist(batch, schema=SCHEMA))
                batch = []
        if batch:
            writer.write_table(pa.Table.from_pylist(batch, schema=SCHEMA))
        writer.close()
        writer = None
        part.rename(path)
        sql_sheet = book["SQL"]
        sql_sheet.reset_dimensions()
        sql = "\n".join(scalar(v) for r in sql_sheet.iter_rows(values_only=True) for v in r if scalar(v))
        query_path = RAW / f"{source['id']}_{family}.sql"
        query_path.write_text(sql + "\n", encoding="utf-8")
        def filters(alias, column):
            return sorted(set(re.findall(rf"\b{alias}\.{column}\s*=\s*'(\d+)'", sql, flags=re.I)))
        info.update(family=family, header=header, **counter, statistics_years=filters("a", "year"),
                    route_description_years=filters("b", "year"), route_description_months=filters("b", "month"),
                    sql_path=str(query_path.relative_to(CASE)), sql_sha256=sha256(query_path),
                    path=str(path.relative_to(CASE)), sha256=sha256(path), source_sha256=source["sha256"])
        info.update(identity)
        write_json(sidecar, info)
        print(f"{source['id']} {family}: {counter['rows']} rows", flush=True)
        return info
    finally:
        if writer is not None:
            writer.close()
        book.close()


def run():
    cfg = read_yaml(CONFIG)
    source = CASE / cfg["source_manifest"]
    manifest = json.loads(source.read_text())
    if manifest.get("failures") or {s["year"] for s in manifest["files"]} != set(cfg["years"]):
        raise ValueError("source collection incomplete")
    target = OUT / "extraction.json"
    expected = dict(configuration_sha256=sha256(CONFIG), source_manifest_sha256=sha256(source), code_sha256=sha256(__file__))
    if target.exists():
        saved = json.loads(target.read_text())
        if any(saved[k] != v for k, v in expected.items()):
            raise ValueError("existing extraction has different inputs")
        for file in saved["files"]:
            if sha256(CASE / file["path"]) != file["sha256"] or sha256(CASE / file["sql_path"]) != file["sql_sha256"]:
                raise ValueError("cached extraction changed")
        print("verified cached extraction", flush=True)
        return
    if shutil.disk_usage(CASE).free < cfg["min_free_gib"] * 1024 ** 3:
        raise ValueError("free disk below configured floor")
    RAW.mkdir(parents=True, exist_ok=True)
    results = []
    for entry in manifest["files"]:
        path = CASE / entry["path"]
        if sha256(path) != entry["sha256"]:
            raise ValueError("source checksum mismatch")
        if path.suffix == ".xlsx":
            results.append(parse_book(entry, "", path, cfg))
        else:
            with zipfile.ZipFile(path) as archive:
                if any(not n.endswith(".xlsx") for n in archive.namelist()):
                    raise ValueError("unexpected nested archive member")
                for member in archive.namelist():
                    results.append(parse_book(entry, member, io.BytesIO(archive.read(member)), cfg))
        write_json(RAW / "extraction_progress.json", {**expected, "files": results})
    write_json(target, {"completed_at": utcnow(), **expected, "configuration": cfg, "files": results})
