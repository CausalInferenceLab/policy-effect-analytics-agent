"""고정 원본을 해시로 검증해 수집하고 버스 ZIP·도로 XLSX를 정규화한다."""

import argparse
import calendar
import csv
import io
import math
from datetime import timedelta
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import numpy as np
import pandas as pd
import requests
from _shared import CASE, checked, config, digest, load_json, new_run, reference, sha256, write_json
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

SOURCE = "https://topis.seoul.go.kr/refRoom/openRefRoom_1.do"
META = [
    "일자",
    "요일",
    "도로명",
    "링크아이디",
    "시점명",
    "종점명",
    "방향",
    "거리",
    "차선수",
    "기능유형구분",
    "도심/외곽구분",
    "권역구분",
]
HEADERS = META + [f"~{i:02d}시" for i in range(1, 25)]
LEGACY_HEADERS = META + [f"{i:02d}시" for i in range(1, 25)]
RENAME = {
    "기준_날짜": "date",
    "노선_ID": "route",
    "출발_정류장_ID": "stop",
    "출발_정류장_순서": "sequence",
    "도착_정류장_ID": "arrival",
    "도착_정류장_순서": "arrival_sequence",
}


def session():
    s = requests.Session()
    # 조회용 POST만 사용한다. 게시·수정 요청은 없다.
    retry = Retry(
        total=3,
        backoff_factor=1,
        status_forcelist=[429, 500, 502, 503, 504],
        allowed_methods={"GET", "POST"},
    )
    s.mount("https://", HTTPAdapter(max_retries=retry))
    s.headers.update({"Referer": SOURCE, "User-Agent": "public-data-course-gangnam/1.0"})
    return s


def validate_xlsx(path):
    try:
        with ZipFile(path) as z:
            if "xl/workbook.xml" not in z.namelist() or z.testzip() is not None:
                raise ValueError("XLSX 구조 또는 CRC 오류")
    except BadZipFile as e:
        raise ValueError("XLSX 대신 HTML 오류 응답 또는 손상된 ZIP을 받았습니다") from e


def interval_start(day, hour_label):
    """공식 종료시각 표기 확인 후 사용. ~24시도 원래 날짜의 23시다."""
    if hour_label not in HEADERS[12:]:
        raise ValueError(f"알 수 없는 시간 열: {hour_label}")
    raw = str(day).strip()
    if raw.endswith(".0"):
        raw = raw[:-2]
    date = pd.to_datetime(raw, format="%Y%m%d", errors="raise")
    return date.tz_localize("Asia/Seoul") + timedelta(hours=int(hour_label[1:3]) - 1)


def speed_value(raw, rules):
    if raw is None or str(raw).strip() == "":
        if not rules["blank_is_missing"]:
            raise ValueError("빈 셀의 처리 정의가 없습니다")
        return math.nan, "blank"
    if str(raw).strip() in {str(x) for x in rules["sentinels"]}:
        return math.nan, "sentinel"
    try:
        value = float(raw)
    except (TypeError, ValueError) as e:
        raise ValueError(f"정의되지 않은 속도 값: {raw!r}") from e
    if not math.isfinite(value):
        raise ValueError("정의되지 않은 비유한 속도 값")
    for marker in rules["sentinels"]:
        try:
            numeric_marker = float(marker)
        except (TypeError, ValueError):
            continue
        if value == numeric_marker:
            return math.nan, "sentinel"
    if value < 0:
        if rules["negative"] != "missing":
            raise ValueError("음수 정의 미확인")
        return math.nan, "negative"
    if value == 0:
        if rules["zero"] == "missing":
            return math.nan, "zero"
        if rules["zero"] != "valid":
            raise ValueError("0 속도 정의 미확인")
    return value, ""


def parse_month(path, month, roads, rules=None):
    """rules=None이면 속도값에 접근하지 않고 메타데이터만 추출한다."""
    from openpyxl import load_workbook

    meta, records, seen, first = {}, [], set(), None
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = list(next(rows))
        if header not in (HEADERS, LEGACY_HEADERS):
            raise ValueError(f"헤더 변경 감지: {header}")
        for row in rows:
            if all(v is None for v in row):
                continue
            if len(row) != len(HEADERS):
                raise ValueError("열 개수 불일치")
            if row[2] not in roads:
                continue
            day = interval_start(row[0], "~01시")
            if day.strftime("%Y-%m") != month:
                raise ValueError(f"요청 월 밖의 자료: {day}")
            if row[3] is None or str(row[3]).strip() == "":
                raise ValueError("링크 ID가 비어 있습니다")
            link = str(row[3]).removesuffix(".0")
            key = (link, str(day.date()))
            if key in seen:
                raise ValueError(f"링크×일자 중복: {key}")
            seen.add(key)
            m = dict(zip(META[2:], row[2:12], strict=True))
            m["링크아이디"] = link
            variant = digest(m)
            meta.setdefault(
                variant,
                {**m, "observed_from": str(day.date()), "observed_to": str(day.date()), "days": 0},
            )
            meta[variant]["observed_to"] = max(meta[variant]["observed_to"], str(day.date()))
            meta[variant]["observed_from"] = min(meta[variant]["observed_from"], str(day.date()))
            meta[variant]["days"] += 1
            first = day if first is None else min(first, day)
            if rules is not None:
                for i, raw in enumerate(row[12:]):
                    value, reason = speed_value(raw, rules)
                    records.append((link, day + timedelta(hours=i), value, not reason, reason))
        if not seen:
            raise ValueError(f"선택 도로의 자료가 없습니다: {month}")
    finally:
        wb.close()
    expected_days = calendar.monthrange(int(month[:4]), int(month[5:]))[1]
    summary = {
        "month": month,
        "header": header,
        "selected_link_days": len(seen),
        "expected_days": expected_days,
        "metadata": sorted(meta.values(), key=lambda x: (x["링크아이디"], x["observed_from"])),
    }
    frame = pd.DataFrame(
        records,
        columns=["link_id", "interval_start_kst", "speed_kmh", "is_valid", "invalid_reason"],
    )
    return frame, summary


def extract_bus(path, stops, hours):
    """열이 어긋난 주간 파일은 격리하고 양수·0·음수·공란을 원형대로 보존한다."""
    keys = list(RENAME)
    columns = [f"운행시간_{hour:02}시" for hour in hours]
    parts, dates, members_audit = [], set(), []
    with ZipFile(path) as archive:
        members = sorted(n for n in archive.namelist() if n.lower().endswith(".csv"))
        if not members:
            raise ValueError("버스 압축파일에 CSV가 없습니다")
        for name in members:
            with archive.open(name) as stream:
                reader = csv.reader(io.TextIOWrapper(stream, encoding="cp949"))
                header, first = next(reader, []), next(reader, None)
                if not set(keys + columns).issubset(header) or (
                    first is not None and len(header) != len(first)
                ):
                    members_audit.append(
                        dict(member=name, status="quarantined_header_row_mismatch")
                    )
                    continue
                if first is None:
                    members_audit.append(dict(member=name, status="empty"))
                    continue
            members_audit.append(dict(member=name, status="accepted"))
            with archive.open(name) as stream:
                for chunk in pd.read_csv(
                    stream, encoding="cp949", dtype=str, usecols=keys + columns, chunksize=100000
                ):
                    pd.to_datetime(chunk["기준_날짜"].unique(), format="%Y%m%d", errors="raise")
                    dates.update(chunk["기준_날짜"].unique())
                    chosen = chunk["출발_정류장_ID"].isin(stops) | chunk["도착_정류장_ID"].isin(
                        stops
                    )
                    frame = chunk.loc[chosen].copy()
                    for column in columns:
                        frame[column] = pd.to_numeric(frame[column], errors="raise")
                        if np.isinf(frame[column]).any():
                            raise ValueError("버스 원본에 무한대 값이 있습니다")
                    parts.append(frame)
    if not any(m["status"] == "accepted" for m in members_audit):
        raise ValueError("사용 가능한 버스 원본 멤버가 없습니다")
    rename = RENAME | dict(zip(columns, [f"h{h:02}" for h in hours], strict=True))
    frame = pd.concat(parts, ignore_index=True).rename(columns=rename)
    if frame[list(RENAME.values())].isna().any().any():
        raise ValueError("버스 결합 키가 비었습니다")
    return frame, dict(dates=sorted(dates), member_validation=members_audit)


def source_file(entry, local_root, folder, client):
    """기존 원본이나 신규 다운로드가 고정 바이트와 다르면 조용히 교체하지 않는다."""
    from _shared import inside

    local = inside(local_root, entry["path"])
    if local.exists():
        return checked(entry, local_root)
    suffix = ".zip" if entry["kind"] == "bus_time" else ".xlsx"
    path = folder / (entry["kind"] + "_" + entry["month"] + suffix)
    temp = path.with_suffix(path.suffix + ".part")
    try:
        if entry["method"] not in {"GET", "POST"} or not entry["url"].startswith("https://"):
            raise ValueError("지원하지 않는 원본 조회 방식")
        options = {"data" if entry["method"] == "POST" else "params": entry["request"]}
        with client.request(
            entry["method"], entry["url"], stream=True, timeout=(15, 180), **options
        ) as response:
            response.raise_for_status()
            with temp.open("wb") as stream:
                for chunk in response.iter_content(1024 * 1024):
                    stream.write(chunk)
        if temp.stat().st_size != entry["bytes"] or sha256(temp) != entry["sha256"]:
            raise ValueError(
                f"공급 원본 버전 변경: {entry['kind']} {entry['month']}. 자동 대체하지 않습니다"
            )
        temp.replace(path)
        return path
    finally:
        temp.unlink(missing_ok=True)


def metadata_exclusions(metadata, cfg):
    result = []
    fields = cfg["link_selection"]["stable_fields"]
    for link, spec in cfg["selected_links"].items():
        rows = [row for row in metadata if row["링크아이디"] == link]
        signatures = [
            {k: float(row[k]) if k == "거리" else str(row[k]) for k in fields} for row in rows
        ]
        if not rows or any(signature != spec["metadata"] for signature in signatures):
            result.append(link)
    return result


def collect(run_id, kind="all", local_root=CASE, reuse_cache=False):
    import shutil

    from _shared import inside

    cfg = config()
    entries = load_json(CASE / "source-manifest.json")["files"]
    selected = entries if kind == "all" else [e for e in entries if e["kind"] == kind]
    if not selected or len({(e["kind"], e["month"]) for e in selected}) != len(selected):
        raise ValueError("비어 있거나 중복된 원본 목록")
    out = new_run(run_id)
    raw = out / "raw"
    raw.mkdir()
    records = []
    stops = sorted({s[k] for s in cfg["bus"]["segments"] for k in ["stop", "arrival"]})
    with session() as client:
        for entry in selected:
            cache = entry["cache"]
            target = out / (entry["kind"] + "_" + entry["month"] + ".parquet")
            if cache["raw_sha256"] != entry["sha256"]:
                raise ValueError("원본과 캐시 계보 불일치")
            if reuse_cache and inside(local_root, cache["path"]).exists():
                shutil.copy2(checked(cache, local_root), target)
                info = {k: cache[k] for k in ["dates", "scopes", "masked_links"] if k in cache}
                info["method"] = "verified_normalized_cache"
            else:
                path = source_file(entry, local_root, raw, client)
                if entry["kind"] == "bus_time":
                    frame, info = extract_bus(path, stops, cfg["bus"]["hours"])
                    info["scopes"] = cache["scopes"]
                else:
                    validate_xlsx(path)
                    frame, summary = parse_month(
                        path,
                        entry["month"],
                        cfg["forecast"]["roads"],
                        cfg["forecast"]["speed_rules"],
                    )
                    masked = metadata_exclusions(summary["metadata"], cfg["forecast"])
                    if masked != cache["masked_links"]:
                        raise ValueError("고정 원본의 링크 속성 대조가 기존 기록과 다릅니다")
                    info = dict(masked_links=masked)
                frame.to_parquet(target, index=False)
                info["method"] = "parsed_raw_source"
            records.append(
                dict(
                    kind=entry["kind"],
                    month=entry["month"],
                    raw_sha256=entry["sha256"],
                    **reference(target),
                    **info,
                )
            )
            print(f"정규화 완료: {entry['kind']} {entry['month']}", flush=True)
    result = dict(inputs=load_json(out / "run.json")["inputs"], files=records)
    write_json(out / "inputs.json", result)
    return out / "inputs.json"


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-id", required=True)
    parser.add_argument("--kind", choices=["all", "bus_time", "road_speed"], default="all")
    parser.add_argument("--local-source-root", type=Path, default=CASE)
    parser.add_argument(
        "--reuse-cache", action="store_true", help="명세에 해시가 고정된 로컬 정규화 파일만 재사용"
    )
    args = parser.parse_args()
    print(collect(args.run_id, args.kind, args.local_source_root, args.reuse_cache))
